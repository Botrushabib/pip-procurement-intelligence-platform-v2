from datetime import datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.services.excel_reader import (
    ExcelReader,
    InvalidWorkbookError,
    RequiredHeadersNotFoundError,
    SourceFileNotFoundError,
    WorksheetNotFoundError,
)


HEADERS = [
    "Suppliers",
    "Date",
    "Reference",
    "Product",
    "To",
    "Qty",
    "Unit",
    "Unit Price",
    "Subtotal",
]


def create_workbook(path: Path, sheet_name: str = "Report", header_row: int = 1) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    for row in range(1, header_row):
        worksheet.cell(row, 1, "Odoo export")
    for column, header in enumerate(HEADERS, start=1):
        worksheet.cell(header_row, column, header)
    worksheet.append(
        [
            "Supplier A",
            datetime(2026, 1, 15, 10, 30),
            "PO-001",
            "Product A",
            "Warehouse A",
            2,
            "Each",
            10.5,
            21.0,
        ]
    )
    workbook.save(path)


def test_valid_xlsx_default_sheet_and_raw_record(tmp_path):
    path = tmp_path / "purchases.xlsx"
    create_workbook(path)

    result = ExcelReader().read(path)

    assert result.source_file == str(path)
    assert result.source_sheet == "Report"
    assert result.header_row == 1
    assert result.total_rows == 1
    assert result.transaction_rows == 1
    assert result.skipped_rows == 0
    assert result.records[0].source_row == 2
    assert result.records[0].supplier_raw == "Supplier A"
    assert result.records[0].product_raw == "Product A"
    assert result.records[0].destination_raw == "Warehouse A"
    assert result.records[0].unit_raw == "Each"
    assert isinstance(result.records[0].transaction_date, datetime)


def test_explicit_worksheet_name(tmp_path):
    path = tmp_path / "purchases.xlsx"
    create_workbook(path, sheet_name="Purchases")

    result = ExcelReader().read(path, "Purchases")

    assert result.source_sheet == "Purchases"
    assert len(result) == 1


def test_missing_source_file(tmp_path):
    with pytest.raises(SourceFileNotFoundError, match="Source XLSX file not found"):
        ExcelReader().read(tmp_path / "missing.xlsx")


def test_invalid_workbook(tmp_path):
    path = tmp_path / "invalid.xlsx"
    path.write_bytes(b"not an xlsx workbook")

    with pytest.raises(InvalidWorkbookError, match="Invalid XLSX workbook"):
        ExcelReader().read(path)


def test_missing_worksheet(tmp_path):
    path = tmp_path / "purchases.xlsx"
    create_workbook(path)

    with pytest.raises(WorksheetNotFoundError, match="Worksheet not found: Missing"):
        ExcelReader().read(path, "Missing")


def test_required_headers_missing(tmp_path):
    path = tmp_path / "purchases.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Report"
    worksheet.append(["Date", "Product"])
    workbook.save(path)

    with pytest.raises(RequiredHeadersNotFoundError, match="Required headers not found"):
        ExcelReader().read(path)


def test_header_scan_normalizes_position_case_and_whitespace(tmp_path):
    path = tmp_path / "purchases.xlsx"
    create_workbook(path, header_row=3)
    workbook = load_workbook(path)
    worksheet = workbook["Report"]
    variants = [
        "  SUPPLIERS  ",
        " dAtE ",
        "Reference",
        "Product",
        "To",
        "QTY",
        "Unit",
        "Unit   Price",
        "Subtotal",
    ]
    for column, value in enumerate(variants, start=1):
        worksheet.cell(3, column, value)
    workbook.save(path)

    result = ExcelReader().read(path)

    assert result.header_row == 3
    assert result.records[0].source_row == 4


def test_non_transaction_rows_are_skipped(tmp_path):
    path = tmp_path / "purchases.xlsx"
    create_workbook(path)
    workbook = load_workbook(path)
    worksheet = workbook["Report"]
    worksheet.append(["Supplier", datetime(2026, 1, 16), "", "Product", "To", 1, "Each", 2, 2])
    worksheet.append(["Supplier", datetime(2026, 1, 17), "PO-002", "Product", "To", "1", "Each", 2, 2])
    workbook.save(path)

    result = ExcelReader().read(path)

    assert result.total_rows == 3
    assert result.transaction_rows == 1
    assert result.skipped_rows == 2


@pytest.mark.parametrize("quantity", [1, 1.5, Decimal("2.5")])
def test_supported_numeric_quantity_types(tmp_path, quantity):
    path = tmp_path / "purchases.xlsx"
    create_workbook(path)
    workbook = load_workbook(path)
    workbook["Report"]["F2"] = quantity
    workbook.save(path)

    assert ExcelReader().read(path).records[0].quantity_raw == quantity


@pytest.mark.parametrize("quantity", [True, "1", None])
def test_unsupported_numeric_quantity_values_are_skipped(tmp_path, quantity):
    path = tmp_path / "purchases.xlsx"
    create_workbook(path)
    workbook = load_workbook(path)
    workbook["Report"]["F2"] = quantity
    workbook.save(path)

    assert len(ExcelReader().read(path)) == 0


@pytest.mark.parametrize("quantity", [1 + 2j, object()])
def test_complex_and_arbitrary_quantity_values_are_rejected(quantity):
    values = {
        "date": datetime(2026, 1, 15),
        "reference": "PO-001",
        "product": "Product A",
        "qty": quantity,
        "unit price": 1,
        "subtotal": 1,
    }

    assert ExcelReader._is_transaction_row(values) is False


def test_numeric_unit_price_and_subtotal_are_preserved(tmp_path):
    path = tmp_path / "purchases.xlsx"
    create_workbook(path)

    record = ExcelReader().read(path).records[0]

    assert record.unit_price_raw == 10.5
    assert record.subtotal_raw == 21.0


def test_input_workbook_is_unchanged_and_reads_are_deterministic(tmp_path):
    path = tmp_path / "purchases.xlsx"
    create_workbook(path)
    before = path.read_bytes()

    first = ExcelReader().read(path)
    second = ExcelReader().read(path)

    assert path.read_bytes() == before
    assert first == second