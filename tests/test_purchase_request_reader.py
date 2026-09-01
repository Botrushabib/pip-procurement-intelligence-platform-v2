from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.services.purchase_request_reader import (
    PurchaseRequestExcelReader,
    QuantityClassification,
    RowClassification,
    SourceSchemaError,
    UnsupportedFileError,
    WorkbookReadError,
)


HEADERS = ["Internal Reference", "Name", "Item description", "Unit", "pr"]


def create_workbook(
    path: Path,
    rows: list[list],
    sheet_name: str = "Sheet1",
    header_row: int = 1,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    for row in range(1, header_row):
        worksheet.cell(row, 1, "PR export")
    for column, header in enumerate(HEADERS, start=1):
        worksheet.cell(header_row, column, header)
    for row in rows:
        worksheet.append(row)
    workbook.save(path)


def test_expected_headers_are_detected(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(path, [["161021", "Beef Tenderloin", "Fresh", "kg", "500"]])

    result = PurchaseRequestExcelReader().read(path)

    assert result.source_sheet == "Sheet1"
    assert result.header_row == 1
    assert result.total_rows == 1


def test_missing_required_header_raises_source_schema_error(tmp_path):
    path = tmp_path / "pr.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Internal Reference", "Name", "Unit", "pr"])
    workbook.save(path)

    with pytest.raises(SourceSchemaError, match="Required headers not found"):
        PurchaseRequestExcelReader().read(path)


def test_numeric_pr_is_classified_as_pure_numeric(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(path, [["161021", "Beef Tenderloin", "Fresh", "kg", "500"]])

    result = PurchaseRequestExcelReader().read(path)

    quantity = result.rows[0].quantity
    assert quantity.classification is QuantityClassification.PURE_NUMERIC
    assert quantity.quantity_value == Decimal("500")
    assert quantity.raw_quantity == "500"


def test_explicit_zero_is_preserved_and_not_null(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(path, [["161021", "Beef Tenderloin", "Fresh", "kg", 0]])

    result = PurchaseRequestExcelReader().read(path)

    quantity = result.rows[0].quantity
    assert quantity.classification is QuantityClassification.EXPLICIT_ZERO
    assert quantity.quantity_value == Decimal("0")
    assert quantity.raw_quantity == "0"


def test_one_box_is_classified_as_quantity_with_packaging(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(path, [["161021", "Beef Tenderloin", "Fresh", "kg", "1 box"]])

    result = PurchaseRequestExcelReader().read(path)

    quantity = result.rows[0].quantity
    assert quantity.classification is QuantityClassification.QUANTITY_WITH_PACKAGING
    assert quantity.raw_quantity == "1 box"
    assert quantity.quantity_token == "1"
    assert quantity.packaging_token == "box"


def test_five_box_is_classified_as_quantity_with_packaging(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(path, [["161021", "Beef Tenderloin", "Fresh", "kg", "5 box"]])

    result = PurchaseRequestExcelReader().read(path)

    quantity = result.rows[0].quantity
    assert quantity.classification is QuantityClassification.QUANTITY_WITH_PACKAGING
    assert quantity.quantity_token == "5"
    assert quantity.packaging_token == "box"


def test_eight_gol_is_classified_as_unknown_quantity_token(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(path, [["161021", "Beef Tenderloin", "Fresh", "kg", "8 gol"]])

    result = PurchaseRequestExcelReader().read(path)

    quantity = result.rows[0].quantity
    assert quantity.classification is QuantityClassification.UNKNOWN_QUANTITY_TOKEN
    assert quantity.raw_quantity == "8 gol"
    assert quantity.quantity_token == "8"
    assert quantity.packaging_token == "gol"


def test_textual_quantity_is_preserved_verbatim(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(path, [["161021", "Beef Tenderloin", "Fresh", "kg", "مطلوب"]])

    result = PurchaseRequestExcelReader().read(path)

    quantity = result.rows[0].quantity
    assert quantity.classification is QuantityClassification.TEXTUAL_QUANTITY
    assert quantity.raw_quantity == "مطلوب"
    assert quantity.quantity_value is None


def test_blank_quantity_is_distinguishable_from_zero(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(path, [["161021", "Beef Tenderloin", "Fresh", "kg", None]])

    result = PurchaseRequestExcelReader().read(path)

    quantity = result.rows[0].quantity
    assert quantity.classification is QuantityClassification.BLANK_QUANTITY
    assert quantity.raw_quantity is None
    assert quantity.quantity_value is None


def test_raw_values_are_preserved_for_product_rows(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(
        path, [["161021", "Beef Tenderloin", "Fresh cut", "kg", "5 box"]]
    )

    result = PurchaseRequestExcelReader().read(path)
    row = result.rows[0]

    assert row.raw_internal_reference == "161021"
    assert row.raw_product == "Beef Tenderloin"
    assert row.raw_item_description == "Fresh cut"
    assert row.raw_unit == "kg"
    assert row.quantity.raw_quantity == "5 box"


def test_category_row_is_not_classified_as_product_row(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(
        path,
        [
            [None, "Food / Beef", None, None, None],
            ["161021", "Beef Tenderloin", "Fresh", "kg", "500"],
        ],
    )

    result = PurchaseRequestExcelReader().read(path)

    assert result.rows[0].row_classification is RowClassification.CATEGORY_ROW
    assert result.rows[1].row_classification is RowClassification.PRODUCT_ROW
    assert result.category_row_count == 1
    assert result.product_row_count == 1
    assert result.product_rows == (result.rows[1],)


def test_internal_reference_name_unit_and_pr_is_product_row(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(path, [["161021", "Beef Tenderloin", "Fresh", "kg", "500"]])

    result = PurchaseRequestExcelReader().read(path)

    assert result.rows[0].row_classification is RowClassification.PRODUCT_ROW


def test_missing_internal_reference_with_name_and_unit_and_pr_is_still_product_row(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(path, [[None, "Beef Tenderloin", "Fresh", "kg", "500"]])

    result = PurchaseRequestExcelReader().read(path)

    assert result.rows[0].row_classification is RowClassification.PRODUCT_ROW


def test_missing_internal_reference_with_blank_unit_and_pr_is_category_row(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(path, [[None, "Section Header", None, None, None]])

    result = PurchaseRequestExcelReader().read(path)

    assert result.rows[0].row_classification is RowClassification.CATEGORY_ROW


def test_food_beef_observed_category_row_is_classified_correctly(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(path, [[None, "Food / Beef", None, None, None]])

    result = PurchaseRequestExcelReader().read(path)

    assert result.rows[0].row_classification is RowClassification.CATEGORY_ROW


def test_food_butter_observed_category_row_is_classified_correctly(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(path, [[None, "Food / Butter", None, None, None]])

    result = PurchaseRequestExcelReader().read(path)

    assert result.rows[0].row_classification is RowClassification.CATEGORY_ROW


def test_completely_blank_row_is_blank_or_format_row(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(path, [[None, None, None, None, None]])

    result = PurchaseRequestExcelReader().read(path)

    assert result.rows[0].row_classification is RowClassification.BLANK_OR_FORMAT_ROW


def test_product_row_with_missing_internal_reference_remains_in_product_rows(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(
        path,
        [
            [None, "Beef Tenderloin", "Fresh", "kg", "500"],
            [None, "Food / Beef", None, None, None],
        ],
    )

    result = PurchaseRequestExcelReader().read(path)

    assert result.product_rows == (result.rows[0],)


def test_product_row_with_missing_internal_reference_preserves_raw_fields(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(path, [[None, "Beef Tenderloin", "Fresh", "kg", "500"]], header_row=2)

    result = PurchaseRequestExcelReader().read(path)
    row = result.rows[0]

    assert row.raw_internal_reference is None
    assert row.raw_product == "Beef Tenderloin"
    assert row.raw_unit == "kg"
    assert row.quantity.raw_quantity == "500"
    assert row.source_row == 3


def test_blank_or_format_row_is_classified_correctly(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(
        path,
        [
            [None, None, None, None, None],
            ["161021", "Beef Tenderloin", "Fresh", "kg", "500"],
        ],
    )

    result = PurchaseRequestExcelReader().read(path)

    assert result.rows[0].row_classification is RowClassification.BLANK_OR_FORMAT_ROW
    assert result.blank_row_count == 1


def test_product_row_provenance_matches_original_worksheet_row(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(
        path,
        [
            ["161021", "Beef Tenderloin", "Fresh", "kg", "500"],
            ["161033", "Butter Block", "Salted", "kg", "300"],
        ],
        header_row=3,
    )

    result = PurchaseRequestExcelReader().read(path)

    assert result.header_row == 3
    assert result.rows[0].source_row == 4
    assert result.rows[1].source_row == 5


def test_multiple_sheets_are_read_independently(tmp_path):
    path = tmp_path / "pr.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "Sheet1"
    for column, header in enumerate(HEADERS, start=1):
        first.cell(1, column, header)
    first.append(["161021", "Beef Tenderloin", "Fresh", "kg", "500"])

    second = workbook.create_sheet("Sheet2")
    for column, header in enumerate(HEADERS, start=1):
        second.cell(1, column, header)
    second.append(["118005", "Milk Full Cream", "", "L", "25"])
    workbook.save(path)

    results = PurchaseRequestExcelReader().read_all_sheets(path)

    assert {result.source_sheet for result in results} == {"Sheet1", "Sheet2"}
    sheet_one = next(r for r in results if r.source_sheet == "Sheet1")
    sheet_two = next(r for r in results if r.source_sheet == "Sheet2")
    assert sheet_one.rows[0].raw_product == "Beef Tenderloin"
    assert sheet_two.rows[0].raw_product == "Milk Full Cream"


def test_item_description_is_optional(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(path, [["161021", "Beef Tenderloin", "", "kg", "500"]])

    result = PurchaseRequestExcelReader().read(path)

    assert not result.rows[0].raw_item_description
    assert result.rows[0].row_classification is RowClassification.PRODUCT_ROW


def test_unit_is_preserved_verbatim(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(
        path,
        [
            ["161021", "Beef Tenderloin", "Fresh", "kg", "500"],
            ["170034", "Case of Water", "Fresh", "Each", "10"],
            ["118005", "Milk Full Cream", "Fresh", "L", "25"],
            ["170099", "Loose Item", "Fresh", "UNIT", "1"],
        ],
    )

    result = PurchaseRequestExcelReader().read(path)

    assert [row.raw_unit for row in result.rows] == ["kg", "Each", "L", "UNIT"]


def test_filename_is_preserved_only_as_provenance(tmp_path):
    path = tmp_path / "PR 22.8.2026 freezer (5)(2).xlsx"
    create_workbook(path, [["161021", "Beef Tenderloin", "Fresh", "kg", "500"]])

    result = PurchaseRequestExcelReader().read(path)
    row = result.rows[0]

    assert result.source_file == str(path)
    assert not hasattr(row, "request_date")
    assert not hasattr(row, "reference")
    assert not hasattr(row, "company_id")
    assert not hasattr(row, "branch_id")
    assert not hasattr(row, "warehouse_id")


def test_reader_performs_no_database_access():
    import inspect

    import app.services.purchase_request_reader as reader_module

    source = inspect.getsource(reader_module)
    assert "sqlalchemy" not in source.lower()
    assert "SessionLocal" not in source
    assert "app.core.database" not in source


def test_repeated_reading_is_deterministic(tmp_path):
    path = tmp_path / "pr.xlsx"
    create_workbook(
        path,
        [
            ["161021", "Beef Tenderloin", "Fresh", "kg", "500"],
            ["161033", "Butter Block", "Salted", "kg", "5 box"],
        ],
    )

    reader = PurchaseRequestExcelReader()
    first_result = reader.read(path)
    second_result = reader.read(path)

    assert first_result == second_result


def test_missing_source_file_raises_workbook_read_error(tmp_path):
    with pytest.raises(WorkbookReadError, match="Source XLSX file not found"):
        PurchaseRequestExcelReader().read(tmp_path / "missing.xlsx")


def test_invalid_workbook_raises_unsupported_file_error(tmp_path):
    path = tmp_path / "invalid.xlsx"
    path.write_bytes(b"not an xlsx workbook")

    with pytest.raises(UnsupportedFileError, match="Invalid XLSX workbook"):
        PurchaseRequestExcelReader().read(path)
