from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from enum import Enum
from pathlib import Path
import re
from typing import Any, Iterator, Sequence
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


# The real PR Excel source contract confirmed by direct file inspection
# (Phase 10B.0.1). No Request Date/Reference/Requester/Location/Status/
# Supplier column exists in this source.
REQUIRED_HEADERS = (
    "Internal Reference",
    "Name",
    "Item description",
    "Unit",
    "pr",
)

# Packaging vocabulary observed and documented in the approved business
# contract. Anything outside this set is an unrecognized token, never
# interpreted or converted.
KNOWN_PACKAGING_TOKENS = frozenset({"box", "بالتة"})

_NUMERIC_PATTERN = re.compile(r"\d+(?:\.\d+)?")
_QUANTITY_WITH_TOKEN_PATTERN = re.compile(r"(\d+(?:\.\d+)?)\s+(\S+)")


class PurchaseRequestReaderError(Exception):
    """Base exception for Purchase Request source-reader failures."""


class WorkbookReadError(PurchaseRequestReaderError):
    """The workbook or requested worksheet could not be read."""


class UnsupportedFileError(PurchaseRequestReaderError):
    """The source file is not a readable XLSX workbook."""


class SourceSchemaError(PurchaseRequestReaderError):
    """The required source columns were not found."""


class InvalidStructureError(PurchaseRequestReaderError):
    """The workbook has no usable structure (e.g. no worksheets)."""


class RowClassification(str, Enum):
    PRODUCT_ROW = "PRODUCT_ROW"
    CATEGORY_ROW = "CATEGORY_ROW"
    BLANK_OR_FORMAT_ROW = "BLANK_OR_FORMAT_ROW"


class QuantityClassification(str, Enum):
    PURE_NUMERIC = "PURE_NUMERIC"
    EXPLICIT_ZERO = "EXPLICIT_ZERO"
    QUANTITY_WITH_PACKAGING = "QUANTITY_WITH_PACKAGING"
    UNKNOWN_QUANTITY_TOKEN = "UNKNOWN_QUANTITY_TOKEN"
    TEXTUAL_QUANTITY = "TEXTUAL_QUANTITY"
    BLANK_QUANTITY = "BLANK_QUANTITY"


@dataclass(frozen=True)
class QuantityReading:
    classification: QuantityClassification
    raw_quantity: str | None
    quantity_value: Decimal | None = None
    quantity_token: str | None = None
    packaging_token: str | None = None


@dataclass(frozen=True)
class PurchaseRequestSourceRow:
    source_file: str
    source_sheet: str
    source_row: int
    row_classification: RowClassification
    raw_internal_reference: Any
    raw_product: Any
    raw_item_description: Any
    raw_unit: Any
    quantity: QuantityReading


@dataclass(frozen=True)
class PurchaseRequestSheetResult(Sequence[PurchaseRequestSourceRow]):
    source_file: str
    source_sheet: str
    header_row: int
    total_rows: int
    product_row_count: int
    category_row_count: int
    blank_row_count: int
    rows: tuple[PurchaseRequestSourceRow, ...]

    def __iter__(self) -> Iterator[PurchaseRequestSourceRow]:
        return iter(self.rows)

    def __len__(self) -> int:
        return len(self.rows)

    def __getitem__(
        self, index: int | slice
    ) -> PurchaseRequestSourceRow | tuple[PurchaseRequestSourceRow, ...]:
        return self.rows[index]

    @property
    def product_rows(self) -> tuple[PurchaseRequestSourceRow, ...]:
        return tuple(
            row for row in self.rows if row.row_classification is RowClassification.PRODUCT_ROW
        )


def _normalize_header(value: Any) -> str:
    if not isinstance(value, str):
        return ""
    return re.sub(r"\s+", " ", value.strip()).lower()


def _is_meaningful(value: Any) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _raw_text(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _classify_quantity(raw_value: Any) -> QuantityReading:
    if not _is_meaningful(raw_value):
        return QuantityReading(QuantityClassification.BLANK_QUANTITY, raw_quantity=None)

    if isinstance(raw_value, (int, float, Decimal)) and not isinstance(raw_value, bool):
        text = _raw_text(raw_value)
        value = Decimal(text)
        classification = (
            QuantityClassification.EXPLICIT_ZERO
            if value == 0
            else QuantityClassification.PURE_NUMERIC
        )
        return QuantityReading(classification, raw_quantity=text, quantity_value=value)

    text = str(raw_value).strip()

    if _NUMERIC_PATTERN.fullmatch(text):
        value = Decimal(text)
        classification = (
            QuantityClassification.EXPLICIT_ZERO
            if value == 0
            else QuantityClassification.PURE_NUMERIC
        )
        return QuantityReading(classification, raw_quantity=text, quantity_value=value)

    match = _QUANTITY_WITH_TOKEN_PATTERN.fullmatch(text)
    if match:
        quantity_token, packaging_token = match.groups()
        classification = (
            QuantityClassification.QUANTITY_WITH_PACKAGING
            if packaging_token.lower() in KNOWN_PACKAGING_TOKENS
            else QuantityClassification.UNKNOWN_QUANTITY_TOKEN
        )
        return QuantityReading(
            classification,
            raw_quantity=text,
            quantity_token=quantity_token,
            packaging_token=packaging_token,
        )

    return QuantityReading(QuantityClassification.TEXTUAL_QUANTITY, raw_quantity=text)


def _classify_row(
    internal_reference: Any, name: Any, unit: Any, quantity_raw: Any
) -> RowClassification:
    if _is_meaningful(internal_reference):
        return RowClassification.PRODUCT_ROW
    if _is_meaningful(name):
        # Name-only rows still count as product rows when a product-line
        # field (Unit or pr) is present; Phase 10B.3 may later resolve the
        # product via Name -> Product.name even without Internal Reference.
        if _is_meaningful(unit) or _is_meaningful(quantity_raw):
            return RowClassification.PRODUCT_ROW
        return RowClassification.CATEGORY_ROW
    return RowClassification.BLANK_OR_FORMAT_ROW


def _find_headers(worksheet: Any) -> tuple[int, dict[str, int]]:
    expected = {_normalize_header(header): header.lower() for header in REQUIRED_HEADERS}
    for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        found: dict[str, int] = {}
        for index, value in enumerate(row):
            normalized = _normalize_header(value)
            if normalized in expected:
                found.setdefault(expected[normalized], index)
        if len(found) == len(expected):
            return row_number, found
    raise SourceSchemaError(
        f"Required headers not found in worksheet '{worksheet.title}': "
        f"{', '.join(REQUIRED_HEADERS)}"
    )


def _read_sheet(
    worksheet: Any,
    *,
    source_file: str,
    source_sheet: str,
) -> PurchaseRequestSheetResult:
    header_row, headers = _find_headers(worksheet)
    rows: list[PurchaseRequestSourceRow] = []
    product_row_count = 0
    category_row_count = 0
    blank_row_count = 0

    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        values = {name: row[index] for name, index in headers.items()}
        classification = _classify_row(
            values["internal reference"], values["name"], values["unit"], values["pr"]
        )

        if classification is RowClassification.PRODUCT_ROW:
            product_row_count += 1
        elif classification is RowClassification.CATEGORY_ROW:
            category_row_count += 1
        else:
            blank_row_count += 1

        rows.append(
            PurchaseRequestSourceRow(
                source_file=source_file,
                source_sheet=source_sheet,
                source_row=row_number,
                row_classification=classification,
                raw_internal_reference=values["internal reference"],
                raw_product=values["name"],
                raw_item_description=values["item description"],
                raw_unit=values["unit"],
                quantity=_classify_quantity(values["pr"]),
            )
        )

    return PurchaseRequestSheetResult(
        source_file=source_file,
        source_sheet=source_sheet,
        header_row=header_row,
        total_rows=len(rows),
        product_row_count=product_row_count,
        category_row_count=category_row_count,
        blank_row_count=blank_row_count,
        rows=tuple(rows),
    )


class PurchaseRequestExcelReader:
    def __init__(self, source_file: str | Path | None = None) -> None:
        self.source_file = source_file

    def read(
        self,
        source_file: str | Path | None = None,
        worksheet_name: str | None = None,
    ) -> PurchaseRequestSheetResult:
        path, workbook = self._open_workbook(source_file)
        try:
            if not workbook.sheetnames:
                raise InvalidStructureError(f"Workbook has no worksheets: {path}")

            selected_sheet = worksheet_name or workbook.sheetnames[0]
            if selected_sheet not in workbook.sheetnames:
                raise InvalidStructureError(f"Worksheet not found: {selected_sheet}")

            worksheet = workbook[selected_sheet]
            return _read_sheet(
                worksheet, source_file=str(path), source_sheet=selected_sheet
            )
        finally:
            workbook.close()

    def read_all_sheets(
        self,
        source_file: str | Path | None = None,
    ) -> tuple[PurchaseRequestSheetResult, ...]:
        path, workbook = self._open_workbook(source_file)
        try:
            if not workbook.sheetnames:
                raise InvalidStructureError(f"Workbook has no worksheets: {path}")

            results = []
            for sheet_name in workbook.sheetnames:
                try:
                    results.append(
                        _read_sheet(
                            workbook[sheet_name],
                            source_file=str(path),
                            source_sheet=sheet_name,
                        )
                    )
                except SourceSchemaError:
                    continue

            if not results:
                raise SourceSchemaError(
                    f"No worksheet in '{path}' contains the required PR headers: "
                    f"{', '.join(REQUIRED_HEADERS)}"
                )
            return tuple(results)
        finally:
            workbook.close()

    def _open_workbook(self, source_file: str | Path | None) -> tuple[Path, Any]:
        selected_file = source_file if source_file is not None else self.source_file
        if selected_file is None:
            raise WorkbookReadError("Source XLSX file was not provided")

        path = Path(selected_file)
        if not path.is_file():
            raise WorkbookReadError(f"Source XLSX file not found: {path}")

        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except (BadZipFile, InvalidFileException, OSError, ValueError, TypeError) as error:
            raise UnsupportedFileError(f"Invalid XLSX workbook: {path}") from error

        return path, workbook


def read_purchase_request_excel(
    source_file: str | Path,
    worksheet_name: str | None = None,
) -> PurchaseRequestSheetResult:
    return PurchaseRequestExcelReader().read(source_file, worksheet_name)
