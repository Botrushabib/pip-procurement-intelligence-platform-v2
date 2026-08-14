from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import re
from typing import Any, Iterator, Sequence
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


DEFAULT_WORKSHEET_NAME = "Report"
REQUIRED_HEADERS = (
    "Suppliers",
    "Date",
    "Reference",
    "Product",
    "To",
    "Qty",
    "Unit",
    "Unit Price",
    "Subtotal",
)


class ExcelReaderError(Exception):
    """Base exception for structural Excel reader failures."""


class SourceFileNotFoundError(ExcelReaderError):
    pass


class InvalidWorkbookError(ExcelReaderError):
    pass


class WorksheetNotFoundError(ExcelReaderError):
    pass


class RequiredHeadersNotFoundError(ExcelReaderError):
    pass


@dataclass(frozen=True)
class RawPurchaseTransactionRecord:
    source_file: str
    source_sheet: str
    source_row: int
    transaction_date: Any
    reference: Any
    supplier_raw: Any
    product_raw: Any
    destination_raw: Any
    quantity_raw: Any
    unit_raw: Any
    unit_price_raw: Any
    subtotal_raw: Any


@dataclass(frozen=True)
class ExcelInspectionResult(Sequence[RawPurchaseTransactionRecord]):
    source_file: str
    source_sheet: str
    header_row: int
    total_rows: int
    transaction_rows: int
    skipped_rows: int
    records: tuple[RawPurchaseTransactionRecord, ...]

    def __iter__(self) -> Iterator[RawPurchaseTransactionRecord]:
        return iter(self.records)

    def __len__(self) -> int:
        return len(self.records)

    def __getitem__(self, index: int | slice) -> RawPurchaseTransactionRecord | tuple[RawPurchaseTransactionRecord, ...]:
        return self.records[index]


def _normalize_header(value: Any) -> str:
    if not isinstance(value, str):
        return ""
    return re.sub(r"\s+", " ", value.strip()).lower()


def _is_numeric(value: Any) -> bool:
    return isinstance(value, (int, float, Decimal)) and not isinstance(value, bool)


def _is_meaningful(value: Any) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


class ExcelReader:
    def __init__(
        self,
        source_file: str | Path | None = None,
        worksheet_name: str = DEFAULT_WORKSHEET_NAME,
    ) -> None:
        self.source_file = source_file
        self.worksheet_name = worksheet_name

    def read(
        self,
        source_file: str | Path | None = None,
        worksheet_name: str | None = None,
    ) -> ExcelInspectionResult:
        selected_file = source_file if source_file is not None else self.source_file
        if selected_file is None:
            raise SourceFileNotFoundError("Source XLSX file was not provided")

        path = Path(selected_file)
        if not path.is_file():
            raise SourceFileNotFoundError(f"Source XLSX file not found: {path}")

        selected_sheet = worksheet_name or self.worksheet_name
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except (BadZipFile, InvalidFileException, OSError, ValueError, TypeError) as error:
            raise InvalidWorkbookError(f"Invalid XLSX workbook: {path}") from error

        try:
            if selected_sheet not in workbook.sheetnames:
                raise WorksheetNotFoundError(
                    f"Worksheet not found: {selected_sheet}"
                )
            worksheet = workbook[selected_sheet]
            header_row, headers = self._find_headers(worksheet)
            records: list[RawPurchaseTransactionRecord] = []
            total_rows = max(0, worksheet.max_row - header_row)

            for row_number, row in enumerate(
                worksheet.iter_rows(min_row=header_row + 1, values_only=True),
                start=header_row + 1,
            ):
                values = {name: row[index] for name, index in headers.items()}
                if not self._is_transaction_row(values):
                    continue
                records.append(
                    RawPurchaseTransactionRecord(
                        source_file=str(path),
                        source_sheet=selected_sheet,
                        source_row=row_number,
                        transaction_date=values["date"],
                        reference=values["reference"],
                        supplier_raw=values["suppliers"],
                        product_raw=values["product"],
                        destination_raw=values["to"],
                        quantity_raw=values["qty"],
                        unit_raw=values["unit"],
                        unit_price_raw=values["unit price"],
                        subtotal_raw=values["subtotal"],
                    )
                    )

            return ExcelInspectionResult(
                source_file=str(path),
                source_sheet=selected_sheet,
                header_row=header_row,
                total_rows=total_rows,
                transaction_rows=len(records),
                skipped_rows=total_rows - len(records),
                records=tuple(records),
            )
        finally:
            workbook.close()

    def _find_headers(self, worksheet: Any) -> tuple[int, dict[str, int]]:
        expected = {_normalize_header(header): header.lower() for header in REQUIRED_HEADERS}
        for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            found = {}
            for index, value in enumerate(row):
                normalized = _normalize_header(value)
                if normalized in expected:
                    found.setdefault(expected[normalized], index)
            if len(found) == len(expected):
                return row_number, found
        raise RequiredHeadersNotFoundError(
            f"Required headers not found in worksheet '{worksheet.title}': "
            f"{', '.join(REQUIRED_HEADERS)}"
        )

    @staticmethod
    def _is_transaction_row(values: dict[str, Any]) -> bool:
        return (
            isinstance(values["date"], datetime)
            and _is_meaningful(values["reference"])
            and _is_meaningful(values["product"])
            and _is_numeric(values["qty"])
            and _is_numeric(values["unit price"])
            and _is_numeric(values["subtotal"])
        )


OdooExcelReader = ExcelReader


def read_odoo_excel(
    source_file: str | Path,
    worksheet_name: str = DEFAULT_WORKSHEET_NAME,
) -> ExcelInspectionResult:
    return ExcelReader().read(source_file, worksheet_name)