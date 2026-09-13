from __future__ import annotations

from dataclasses import dataclass
from enum import Enum
from pathlib import Path
import re
from typing import Any, Iterator, Sequence
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


REQUIRED_HEADERS = (
    "Suppliers",
    "Date",
    "Reference",
    "Product",
    "To",
    "Qty",
    "Unit",
    "Unit Price",
    "Subtotal",
)


class ReceivingReaderError(Exception):
    """Base exception for receiving workbook reader failures."""


class SourceFileNotFoundError(ReceivingReaderError):
    pass


class InvalidWorkbookError(ReceivingReaderError):
    pass


class WorksheetNotFoundError(ReceivingReaderError):
    pass


class RequiredHeadersNotFoundError(ReceivingReaderError):
    pass


class ReceivingRowClassification(str, Enum):
    SUPPLIER_HEADER = "SUPPLIER_HEADER"
    PRODUCT_SUMMARY = "PRODUCT_SUMMARY"
    DETAIL = "DETAIL"
    OTHER = "OTHER"


@dataclass(frozen=True)
class ReceivingSourceRow:
    source_file: str
    source_sheet: str
    source_row: int
    row_classification: ReceivingRowClassification
    raw_supplier: Any
    raw_date_1: Any
    raw_date_2: Any
    event_date_candidate: Any
    raw_reference: Any
    raw_product: Any
    product_code: str | None
    raw_item_description: Any
    raw_destination: Any
    raw_quantity: Any
    raw_unit: Any
    raw_unit_price: Any
    raw_subtotal: Any


@dataclass(frozen=True)
class ReceivingSheetResult(Sequence[ReceivingSourceRow]):
    source_file: str
    source_sheet: str
    header_row: int
    total_rows: int
    detail_row_count: int
    non_detail_row_count: int
    date_headers: tuple[str, ...]
    rows: tuple[ReceivingSourceRow, ...]

    def __iter__(self) -> Iterator[ReceivingSourceRow]:
        return iter(self.rows)

    def __len__(self) -> int:
        return len(self.rows)

    def __getitem__(
        self, index: int | slice
    ) -> ReceivingSourceRow | tuple[ReceivingSourceRow, ...]:
        return self.rows[index]

    @property
    def detail_rows(self) -> tuple[ReceivingSourceRow, ...]:
        return tuple(
            row
            for row in self.rows
            if row.row_classification is ReceivingRowClassification.DETAIL
        )


def _normalize_header(value: Any) -> str:
    if not isinstance(value, str):
        return ""
    return re.sub(r"\s+", " ", value.strip()).lower()


def _is_meaningful(value: Any) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _is_numeric(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def extract_product_code(raw_product: Any) -> str | None:
    """Extract a digits-only Odoo code from the approved bracketed format."""
    if not isinstance(raw_product, str):
        return None
    match = re.match(r"^\[(\d+)\](?:\s|$)", raw_product)
    return match.group(1) if match else None


def _find_headers(
    worksheet: Any,
) -> tuple[int, dict[str, int], tuple[str, ...], tuple[int, ...]]:
    required = {_normalize_header(header): header.lower() for header in REQUIRED_HEADERS}
    optional = {"item description": "item description", "description": "item description"}
    expected = {**required, **optional}
    for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        found: dict[str, int] = {}
        date_headers: list[str] = []
        date_indices: list[int] = []
        for index, value in enumerate(row):
            normalized = _normalize_header(value)
            if normalized in expected:
                found.setdefault(expected[normalized], index)
            if "date" in normalized and normalized not in date_headers:
                date_headers.append(normalized)
                date_indices.append(index)
        if len(set(found).intersection(required)) == len(required):
            if "date" not in date_headers:
                date_headers.insert(0, "date")
                date_indices.insert(0, found["date"])
            return row_number, found, tuple(date_headers[:2]), tuple(date_indices[:2])
    raise RequiredHeadersNotFoundError(
        f"Required headers not found in worksheet '{worksheet.title}': "
        f"{', '.join(REQUIRED_HEADERS)}"
    )


def _detail_row(values: dict[str, Any]) -> bool:
    return (
        all(_is_meaningful(values[field]) for field in ("reference", "product", "to", "unit"))
        and all(_is_numeric(values[field]) for field in ("qty", "unit price", "subtotal"))
    )


def _classify_row(values: dict[str, Any]) -> ReceivingRowClassification:
    if _detail_row(values):
        return ReceivingRowClassification.DETAIL
    if (
        _is_meaningful(values["suppliers"])
        and not _is_meaningful(values["reference"])
        and not _is_meaningful(values["product"])
    ):
        return ReceivingRowClassification.SUPPLIER_HEADER
    if _is_meaningful(values["product"]):
        return ReceivingRowClassification.PRODUCT_SUMMARY
    return ReceivingRowClassification.OTHER


class ReceivingExcelReader:
    def __init__(
        self,
        source_file: str | Path | None = None,
        worksheet_name: str = "Report",
    ) -> None:
        self.source_file = source_file
        self.worksheet_name = worksheet_name

    def read(
        self,
        source_file: str | Path | None = None,
        worksheet_name: str | None = None,
    ) -> ReceivingSheetResult:
        selected_file = source_file if source_file is not None else self.source_file
        if selected_file is None:
            raise SourceFileNotFoundError("Source XLSX file was not provided")

        path = Path(selected_file)
        if not path.is_file():
            raise SourceFileNotFoundError(f"Source XLSX file not found: {path}")

        selected_sheet = worksheet_name or self.worksheet_name
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except (BadZipFile, InvalidFileException, OSError, ValueError, TypeError) as error:
            raise InvalidWorkbookError(f"Invalid XLSX workbook: {path}") from error

        try:
            if selected_sheet not in workbook.sheetnames:
                raise WorksheetNotFoundError(f"Worksheet not found: {selected_sheet}")
            worksheet = workbook[selected_sheet]
            header_row, headers, date_headers, date_indices = _find_headers(worksheet)
            rows: list[ReceivingSourceRow] = []
            current_supplier: Any = None
            total_rows = max(0, worksheet.max_row - header_row)

            for row_number, row in enumerate(
                worksheet.iter_rows(min_row=header_row + 1, values_only=True),
                start=header_row + 1,
            ):
                values = {
                    name: row[index] if index < len(row) else None
                    for name, index in headers.items()
                }
                classification = _classify_row(values)
                if classification is ReceivingRowClassification.SUPPLIER_HEADER:
                    current_supplier = values["suppliers"]

                date_values = []
                for index in date_indices:
                    date_values.append(row[index] if index < len(row) else None)
                date_values.extend([None, None])

                rows.append(
                    ReceivingSourceRow(
                        source_file=str(path),
                        source_sheet=selected_sheet,
                        source_row=row_number,
                        row_classification=classification,
                        raw_supplier=(
                            values["suppliers"]
                            if classification is ReceivingRowClassification.SUPPLIER_HEADER
                            else current_supplier
                        ),
                        raw_date_1=date_values[0],
                        raw_date_2=date_values[1],
                        event_date_candidate=values["date"],
                        raw_reference=values["reference"],
                        raw_product=values["product"],
                        product_code=extract_product_code(values["product"]),
                        raw_item_description=values.get("item description"),
                        raw_destination=values["to"],
                        raw_quantity=values["qty"],
                        raw_unit=values["unit"],
                        raw_unit_price=values["unit price"],
                        raw_subtotal=values["subtotal"],
                    )
                )

            detail_row_count = sum(
                row.row_classification is ReceivingRowClassification.DETAIL for row in rows
            )
            return ReceivingSheetResult(
                source_file=str(path),
                source_sheet=selected_sheet,
                header_row=header_row,
                total_rows=total_rows,
                detail_row_count=detail_row_count,
                non_detail_row_count=total_rows - detail_row_count,
                date_headers=date_headers,
                rows=tuple(rows),
            )
        finally:
            workbook.close()


ReceivingEventReader = ReceivingExcelReader